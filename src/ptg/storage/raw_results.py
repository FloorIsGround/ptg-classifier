from __future__ import annotations

import os
import shutil
import tempfile
from dataclasses import dataclass
from pathlib import Path

import openpyxl
import pandas as pd


@dataclass(frozen=True)
class RawResult:
    row_index: int
    timing: str
    speaker: str
    transcript: str
    label: str
    raw_response: str
    parsed: str  # "yes" | "no" | "" (retries exhausted / embed failed)


def write_raw_results(
    workbook_path: str,
    sheet_name: str,
    records: list[RawResult],
) -> None:
    """Write records to a sheet in the raw results workbook.

    Creates the workbook if it doesn't exist. Replaces the sheet if it already
    exists (e.g. a re-run of an aborted infer_sheet call).
    """
    rows = [
        {
            "row_index": r.row_index,
            "timing": r.timing,
            "speaker": r.speaker,
            "transcript": r.transcript,
            "label": r.label,
            "raw_response": r.raw_response,
            "parsed": r.parsed,
        }
        for r in records
    ]
    df = pd.DataFrame(rows)

    target = Path(workbook_path)
    target.parent.mkdir(parents=True, exist_ok=True)
    fd, temp_name = tempfile.mkstemp(prefix=f".{target.stem}.", suffix=".xlsx", dir=target.parent)
    os.close(fd)
    temp_path = Path(temp_name)
    try:
        if target.exists():
            shutil.copy2(target, temp_path)
            mode, extra = "a", {"if_sheet_exists": "replace"}
        else:
            mode, extra = "w", {}
        with pd.ExcelWriter(temp_path, engine="openpyxl", mode=mode, **extra) as writer:
            df.to_excel(writer, sheet_name=sheet_name, index=False)
        os.replace(temp_path, target)
    finally:
        temp_path.unlink(missing_ok=True)


def next_run_sheet_name(workbook_path: str, source_sheet: str) -> str:
    """Return the sheet name for the next run, e.g. 'Team 1 Time 1 (F25) Run 0'."""
    if not os.path.exists(workbook_path):
        return f"{source_sheet} Run 0"
    wb = openpyxl.load_workbook(workbook_path, read_only=True)
    prefix = f"{source_sheet} Run "
    ids: list[int] = []
    for name in wb.sheetnames:
        if name.startswith(prefix):
            try:
                ids.append(int(name[len(prefix):]))
            except ValueError:
                pass
    wb.close()
    next_id = max(ids) + 1 if ids else 0
    return f"{source_sheet} Run {next_id}"


def read_raw_results(workbook_path: str, source_sheet: str) -> pd.DataFrame:
    """Read all run sheets for source_sheet and concatenate into one DataFrame.

    Adds a ``run_id`` column (int) derived from the sheet name suffix.
    Returns an empty DataFrame if no matching sheets are found.
    """
    wb = openpyxl.load_workbook(workbook_path, read_only=True)
    prefix = f"{source_sheet} Run "
    run_sheets: list[tuple[str, int]] = []
    for name in wb.sheetnames:
        if name.startswith(prefix):
            try:
                run_sheets.append((name, int(name[len(prefix):])))
            except ValueError:
                pass
    wb.close()

    if not run_sheets:
        return pd.DataFrame()

    frames: list[pd.DataFrame] = []
    for sheet_name, run_id in sorted(run_sheets, key=lambda x: x[1]):
        df = pd.read_excel(workbook_path, sheet_name=sheet_name)
        df["run_id"] = run_id
        frames.append(df)

    return pd.concat(frames, ignore_index=True)
